"""Filtered, presentation-ready module reports with cell-linked native Excel charts."""

from collections import Counter, defaultdict

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill

from .capacity import capacity_analysis
from .excel_export import BLUE, INK, MUTED, PERCENT_FORMAT, _add_chart, _data_sheet, _excel_value
from .filters import filter_capacity_rows, filter_students
from .models import Attendance, Department, Major, Performance, Room, Shift, Student, Teacher
from .services import matched_scatter


def _report_shell(title, selected, generated_at):
    book = Workbook()
    overview = book.active
    overview.title = "Overview"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:Q2")
    overview["A1"] = f"UNIVERSITY DATA ANALYTICS — {title.upper()}"
    overview["A1"].font = Font(name="Calibri", size=17, bold=True, color="FFFFFF")
    overview["A1"].alignment = Alignment(vertical="center")
    for row in overview["A1:Q2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=INK)
    overview.row_dimensions[1].height = 26
    overview.row_dimensions[2].height = 18
    overview["A3"] = "Report generated"
    overview["C3"] = generated_at.replace(tzinfo=None)
    overview["C3"].number_format = "dd mmm yyyy, hh:mm"
    department = Department.objects.filter(pk=selected["department"]).first() if selected["department"] else None
    major = Major.objects.filter(pk=selected["major"]).first() if selected["major"] else None
    scope = ", ".join([f"Department: {department.code}" if department else "All departments",
                       f"Major: {major.name}" if major else "All majors",
                       f"Shift: {selected['shift']}" if selected["shift"] else "All shifts"])
    overview["A4"] = "Report scope"
    overview["C4"] = scope
    overview["C4"].font = Font(color=MUTED, italic=True)
    overview.column_dimensions["A"].width = 35
    overview.column_dimensions["B"].width = 12
    overview.column_dimensions["C"].width = 36
    for letter in "DEFGHIJKLMNOPQ":
        overview.column_dimensions[letter].width = 13
    overview.freeze_panes = "A6"
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    overview.page_setup.fitToWidth = 1
    overview.page_setup.fitToHeight = 0
    overview.page_setup.orientation = "landscape"
    overview.page_setup.paperSize = overview.PAPERSIZE_A3
    overview.print_options.horizontalCentered = True
    helper = book.create_sheet("_ChartData")
    return book, overview, helper


def _kpis(sheet, entries):
    for row, (label, value, format_code) in enumerate(entries, 6):
        sheet.cell(row, 1, label).font = Font(name="Calibri", bold=True, color=MUTED)
        cell = sheet.cell(row, 3, value if value is not None else "No Data")
        cell.font = Font(name="Calibri", size=12, bold=True, color=INK)
        if isinstance(value, str) and len(value) > 30:
            sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 32
        if value is not None and format_code:
            cell.number_format = format_code


def _sheet(book, title, headers, rows, empty, percentage_columns=(), decimal_columns=()):
    sheet = book.create_sheet(title)
    _data_sheet(sheet, headers, rows, empty, percentage_columns, decimal_columns)
    return sheet


def _insights(book, rows):
    _sheet(book, "Insights", ["Priority", "Area", "Finding / Decision Support"], rows,
           "No actionable findings in the current report scope")


def _comparison(overview, helper, index, title, labels, first_title, first_values, second_title, second_values):
    if not labels:
        return
    column = 1 + index * 4
    for offset, heading in enumerate(("Group", first_title, second_title)):
        helper.cell(1, column + offset, heading)
    for row, (label, first, second) in enumerate(zip(labels, first_values, second_values), 2):
        helper.cell(row, column, _excel_value(label))
        helper.cell(row, column + 1, first)
        helper.cell(row, column + 2, second)
    chart = BarChart()
    chart.title = title
    chart.type = "col"
    chart.grouping = "clustered"
    chart.width, chart.height = 16, 8.5
    chart.visible_cells_only = False
    chart.add_data(Reference(helper, min_col=column + 1, max_col=column + 2, min_row=1, max_row=len(labels) + 1), titles_from_data=True)
    chart.set_categories(Reference(helper, min_col=column, min_row=2, max_row=len(labels) + 1))
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[1].graphicalProperties.solidFill = "94A3B8"
    overview.add_chart(chart, f'{"A" if index % 2 == 0 else "J"}{21 + (index // 2) * 27}')


def _scatter(overview, helper, index, points):
    if not points:
        return
    column = 1 + index * 4
    helper.cell(1, column, "Attendance %")
    helper.cell(1, column + 1, "Score")
    for row, point in enumerate(points, 2):
        helper.cell(row, column, point["x"] / 100)
        helper.cell(row, column + 1, point["y"])
        helper.cell(row, column).number_format = PERCENT_FORMAT
    chart = ScatterChart()
    chart.title = "Attendance vs Score — matched student and subject"
    chart.x_axis.title = "Attendance %"
    chart.y_axis.title = "Score"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 0, 1.05
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 105
    chart.width, chart.height = 16, 9
    chart.visible_cells_only = False
    xvalues = Reference(helper, min_col=column, min_row=2, max_row=len(points) + 1)
    yvalues = Reference(helper, min_col=column + 1, min_row=2, max_row=len(points) + 1)
    series = Series(yvalues, xvalues, title="Matched records")
    series.marker.symbol = "circle"
    series.marker.size = 5
    series.marker.graphicalProperties.solidFill = BLUE
    series.marker.graphicalProperties.line.solidFill = BLUE
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    overview.add_chart(chart, f'{"A" if index % 2 == 0 else "J"}{21 + (index // 2) * 27}')


def _enrollment(book, overview, helper, selected):
    analysis = capacity_analysis()
    groups = filter_capacity_rows(analysis["groups"], selected)
    students = list(filter_students(Student.objects.select_related("department", "major"), selected).order_by("student_id"))
    group_lookup = {(row["major_id"], row["shift"]): row for row in groups}
    _sheet(book, "Enrollment Data", ["Student ID", "Student Name", "Department", "Major", "Shift", "Required Class Groups", "Capacity Status"],
           ([s.student_id, s.student_name, s.department.code, s.major.name, s.shift,
             group_lookup[(s.major_id, s.shift)]["required"], group_lookup[(s.major_id, s.shift)]["status"]] for s in students),
           "No enrollment data available in this scope")
    _sheet(book, "Major & Shift Analysis", ["Department", "Major", "Shift", "Students", "Required Class Groups", "Current Group Seats", "Remaining Current Group Seats"],
           ([g["major__department__code"], g["major__name"], g["shift"], g["students"], g["required"],
             g["current_supported_capacity"], g["remaining_seats"]] for g in groups),
           "No major/shift groups available in this scope")
    pairs = {(g["major__department__code"], g["shift"]) for g in groups}
    teacher_pools = [p for p in analysis["teacher_pools"] if (p["department"], p["shift"]) in pairs]
    shifts = {g["shift"] for g in groups}
    room_pools = [p for p in analysis["shift_pools"] if p["shift"] in shifts]
    _sheet(book, "Capacity Analysis", ["Department", "Major", "Shift", "Students", "Required Class Groups", "Department Shift Demand (Pool)", "Active Department Teachers (Pool)", "Required Rooms", "Shared Shift Room Demand (Pool)", "Active Shared Rooms (Pool)", "Supported Remaining Seats", "Status", "Recommended Shift", "Recommended Supported Seats"],
           ([g["major__department__code"], g["major__name"], g["shift"], g["students"], g["required"],
             g["department_required"], g["teacher_capacity"], g["required_rooms"], g["shift_rooms_used"],
             g["shift_rooms_available"], g["supported_remaining_seats"], g["status"],
             g["recommended_shift"] or "", g["recommended_capacity"]] for g in groups),
           "No capacity analysis available in this scope")
    overview["A16"] = "Resource note"
    overview.merge_cells("C16:H17")
    overview["C16"] = "Teacher demand is department-wide; rooms are shared university-wide by shift. No assignments are implied."
    overview["C16"].alignment = Alignment(wrap_text=True)
    overview.row_dimensions[16].height = 25
    overview.row_dimensions[17].height = 25
    total_required = sum(g["required"] for g in groups)
    _kpis(overview, [("Total Students", len(students), None),
        ("Departments in Scope", len({s.department_id for s in students}), None),
        ("Majors in Scope", len({s.major_id for s in students}), None),
        ("Active Teachers (University)", Teacher.objects.filter(active=True).count(), None),
        ("Active Rooms (Shared)", Room.objects.filter(active=True).count(), None),
        ("Required Class Groups", total_required, None),
        ("Groups Near Capacity", sum(g["status"] == "Near Capacity" for g in groups), None),
        ("Groups at Capacity / Full", sum(g["status"] == "Full" for g in groups), None),
        ("Teacher Shortage Pools", sum(p["shortage"] > 0 for p in teacher_pools), None),
        ("Shift Room Capacity Status", ", ".join(f'{p["shift"]}: {p["status"]}' for p in room_pools) if room_pools else None, None)])
    dept_counts, major_counts, shift_counts = Counter(), Counter(), Counter()
    for s in students:
        dept_counts[s.department.code] += 1
        major_counts[s.major.name] += 1
        shift_counts[s.shift] += 1
    specs = [
        ("Enrollment by Department", dept_counts), ("Enrollment by Major", major_counts),
        ("Enrollment by Shift", {shift: shift_counts[shift] for shift in Shift.values if shift_counts[shift]}),
        ("Required Class Groups by Major & Shift", {f'{g["major__name"]} · {g["shift"]}': g["required"] for g in groups}),
    ]
    for index, (title, counts) in enumerate(specs):
        _add_chart(overview, helper, index, title, list(counts), list(counts.values()), horizontal=index in {1, 3})
    _comparison(overview, helper, 4, "Department / Shift: Required Class Groups vs Active Teachers",
                [f'{p["department"]} · {p["shift"]}' for p in teacher_pools], "Required class groups",
                [p["required"] for p in teacher_pools], "Active teachers", [p["capacity"] for p in teacher_pools])
    _comparison(overview, helper, 5, "Shared Room Demand by Shift (University-wide)",
                [p["shift"] for p in room_pools], "Required rooms", [p["required"] for p in room_pools],
                "Active shared rooms", [p["capacity"] for p in room_pools])
    insights = []
    for p in teacher_pools:
        if p["shortage"]:
            insights.append(["Critical", "Teacher Capacity", f'{p["department"]} {p["shift"]}: {p["required"]} class groups require {p["required"]} teachers; {p["capacity"]} active teachers available. Shortage: {p["shortage"]}.'])
    for p in room_pools:
        if p["shortage"]:
            insights.append(["Warning", "Shared Rooms", f'{p["shift"]}: {p["required"]} simultaneous class groups exceed {p["capacity"]} active shared rooms by {p["shortage"]}.'])
    for g in groups:
        if g["status"] != "Available":
            finding = f'{g["major__name"]} {g["shift"]}: {g["students"]} students, {g["required"]} required class groups, {g["status"]}. '
            finding += f'{g["recommended_shift"]} offers {g["recommended_capacity"]} supported seats for the same major.' if g["recommended_shift"] else "No alternative shift currently has sufficient supported capacity."
            insights.append(["Warning", "Enrollment Availability", finding])
    _insights(book, insights)


def _performance(book, overview, helper, selected):
    records = list(filter_students(Performance.objects.select_related("student__department", "student__major", "subject"), selected, "student__").order_by("student__student_id", "subject__code"))
    _sheet(book, "Student Performance", ["Student ID", "Student Name", "Department", "Major", "Subject Code", "Subject Name", "Score", "Grade", "Pass/Fail"],
           ([r.student.student_id, r.student.student_name, r.student.department.code, r.student.major.name,
             r.subject.code, r.subject.name, float(r.score), r.letter_grade, "Pass" if r.passed else "Fail"] for r in records),
           "No performance data available in this scope", decimal_columns=(7,))
    subjects, majors = defaultdict(list), defaultdict(list)
    grades = Counter()
    for record in records:
        subjects[(record.subject.code, record.subject.name)].append(record)
        majors[record.student.major.name].append(record)
        grades[record.letter_grade] += 1
    subject_rows = [(code, name, len(values), sum(float(v.score) for v in values) / len(values),
                     sum(v.passed for v in values) / len(values)) for (code, name), values in sorted(subjects.items())]
    _sheet(book, "Subject Analysis", ["Subject Code", "Subject", "Records", "Average Score", "Pass Rate"],
           subject_rows, "No subject analysis available", percentage_columns=(5,), decimal_columns=(4,))
    grade_order = ["A", "B+", "B", "C+", "C", "D", "F"]
    _sheet(book, "Grade Distribution", ["Grade", "Records"],
           ([grade, grades[grade]] for grade in grade_order), "No grades available")
    total = len(records)
    passed = sum(record.passed for record in records)
    average = sum(float(record.score) for record in records) / total if total else None
    best = max(subject_rows, key=lambda row: row[3]) if subject_rows else None
    weakest = min(subject_rows, key=lambda row: row[3]) if subject_rows else None
    _kpis(overview, [("Total Performance Records", total, None),
        ("Students Assessed", len({r.student_id for r in records}), None),
        ("Average Score", average, "0.0"), ("Pass Rate", passed / total if total else None, PERCENT_FORMAT),
        ("Fail Rate", (total - passed) / total if total else None, PERCENT_FORMAT),
        ("Highest Average Subject", best[1] if best else None, None),
        ("Lowest Average Subject", weakest[1] if weakest else None, None),
        ("Grade Distribution", ", ".join(f"{grade}: {grades[grade]}" for grade in grade_order if grades[grade]) if total else None, None)])
    _add_chart(overview, helper, 0, "Average Score by Subject", [r[1] for r in subject_rows], [r[3] for r in subject_rows], horizontal=True)
    _add_chart(overview, helper, 1, "Pass vs Fail", ["Pass", "Fail"] if total else [], [passed, total - passed] if total else [], kind="pie", colors=["16A34A", "DC2626"])
    _add_chart(overview, helper, 2, "Grade Distribution", [grade for grade in grade_order if grades[grade]], [grades[grade] for grade in grade_order if grades[grade]])
    _add_chart(overview, helper, 3, "Average Score by Major", list(majors),
               [sum(float(v.score) for v in values) / len(values) for values in majors.values()])
    _add_chart(overview, helper, 4, "Pass Rate by Subject", [r[1] for r in subject_rows], [r[4] * 100 for r in subject_rows], horizontal=True)
    insights = []
    if weakest:
        insights.append(["Warning", "Subject Performance", f'{weakest[1]} has the lowest average score ({weakest[3]:.1f}) among subjects in this report.'])
    if best:
        insights.append(["Info", "Subject Performance", f'{best[1]} has the highest average score ({best[3]:.1f}) among subjects in this report.'])
    if total - passed:
        insights.append(["Warning", "Pass / Fail", f'{total - passed} of {total} assessed records are failing ({(total - passed) / total:.1%}).'])
    if subject_rows:
        top_pass = max(subject_rows, key=lambda row: row[4])
        insights.append(["Info", "Pass / Fail", f'{top_pass[1]} has the highest pass rate ({top_pass[4]:.1%}) in this report.'])
    _insights(book, insights)


def _attendance(book, overview, helper, selected):
    records = list(filter_students(Attendance.objects.select_related("student__major", "student__department", "subject"), selected, "student__").order_by("student__student_id", "subject__code"))
    _sheet(book, "Attendance Records", ["Student ID", "Student Name", "Department", "Major", "Subject Code", "Subject Name", "Total Sessions", "Absent Count", "Attendance %", "Attendance Status"],
           ([r.student.student_id, r.student.student_name, r.student.department.code, r.student.major.name,
             r.subject.code, r.subject.name, r.total_sessions, r.absent_count,
             float(r.attendance_percentage) / 100, r.status] for r in records),
           "No attendance data available in this scope", percentage_columns=(9,))
    subjects, majors, statuses = defaultdict(list), defaultdict(list), Counter()
    for record in records:
        subjects[(record.subject.code, record.subject.name)].append(record)
        majors[record.student.major.name].append(record)
        statuses[record.status] += 1
    subject_rows = [(code, name, len(values), sum(float(v.attendance_percentage) for v in values) / len(values) / 100,
                     sum(v.absent_count for v in values)) for (code, name), values in sorted(subjects.items())]
    _sheet(book, "Subject Attendance", ["Subject Code", "Subject", "Records", "Average Attendance %", "Total Absences"],
           subject_rows, "No subject attendance available", percentage_columns=(4,))
    status_order = ["Good", "Warning", "High Risk", "Critical"]
    _sheet(book, "Risk Analysis", ["Status", "Attendance Records", "Distinct Students"],
           ([status, statuses[status], len({r.student_id for r in records if r.status == status})] for status in status_order),
           "No attendance risk data available")
    performance = filter_students(Performance.objects.select_related("student", "subject"), selected, "student__")
    points, low_scores, good_scores = matched_scatter(performance, records)
    _sheet(book, "Attendance vs Performance", ["Student ID", "Student Name", "Subject", "Attendance %", "Score", "Attendance Status"],
           ([point["student_id"], point["student_name"], point["subject"], point["x"] / 100, point["y"], point["status"]] for point in points),
           "No matching attendance and performance records for the same student and subject", percentage_columns=(4,), decimal_columns=(5,))
    total = len(records)
    average_attendance = sum(float(r.attendance_percentage) for r in records) / total / 100 if total else None
    average_absences = sum(r.absent_count for r in records) / total if total else None
    _kpis(overview, [("Students with Attendance Data", len({r.student_id for r in records}), None),
        ("Attendance Records", total, None), ("Average Attendance %", average_attendance, PERCENT_FORMAT),
        ("Good Attendance Count", statuses["Good"], None), ("Warning Count", statuses["Warning"], None),
        ("High Risk Count", statuses["High Risk"], None), ("Critical Count", statuses["Critical"], None),
        ("Average Absence Count", average_absences, "0.0")])
    _add_chart(overview, helper, 0, "Attendance Risk Distribution", [status for status in status_order if statuses[status]],
               [statuses[status] for status in status_order if statuses[status]], kind="pie",
               colors=[{"Good": "16A34A", "Warning": "D97706", "High Risk": "EA580C", "Critical": "DC2626"}[status] for status in status_order if statuses[status]])
    _add_chart(overview, helper, 1, "Average Attendance by Subject", [r[1] for r in subject_rows], [r[3] * 100 for r in subject_rows], horizontal=True)
    _add_chart(overview, helper, 2, "Average Attendance by Major", list(majors),
               [sum(float(v.attendance_percentage) for v in values) / len(values) for values in majors.values()])
    _add_chart(overview, helper, 3, "Absence Count by Subject", [r[1] for r in subject_rows], [r[4] for r in subject_rows], horizontal=True)
    _scatter(overview, helper, 4, points)
    insights = []
    critical_students = len({r.student_id for r in records if r.status == "Critical"})
    if critical_students:
        insights.append(["Critical", "Attendance Risk", f'{critical_students} students have at least one Critical attendance record in this scope.'])
    if subject_rows:
        weakest = min(subject_rows, key=lambda row: row[3])
        insights.append(["Warning", "Subject Attendance", f'{weakest[1]} has the lowest average attendance ({weakest[3]:.1%}).'])
    if low_scores and good_scores and sum(low_scores) / len(low_scores) < sum(good_scores) / len(good_scores):
        insights.append(["Info", "Attendance vs Performance", f'Matched low-attendance records average {sum(low_scores)/len(low_scores):.1f} score versus {sum(good_scores)/len(good_scores):.1f} for good attendance; this is descriptive, not causal.'])
    _insights(book, insights)


def build_module_report(dataset_type, selected, generated_at=None):
    generated_at = generated_at or timezone.localtime()
    titles = {"enrollment": "Enrollment & Capacity", "performance": "Academic Performance", "attendance": "Attendance Analytics"}
    if dataset_type not in titles:
        raise ValueError("Unknown analytics report")
    book, overview, helper = _report_shell(titles[dataset_type], selected, generated_at)
    {"enrollment": _enrollment, "performance": _performance, "attendance": _attendance}[dataset_type](book, overview, helper, selected)
    book.move_sheet(helper, offset=len(book.sheetnames))
    helper.sheet_state = "hidden"
    return book
