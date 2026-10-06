"""Synthetic demo hierarchy, student names, and cross-file identity checks."""

import os
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import load_workbook

from .demo_names import cambodian_student_names
from .demo_workbooks import MASTER_ROWS, demo_workbooks, messy_demo_workbooks
from .models import Department, Major, Student
from .services import import_rows, validate_workbook
from .setup_import import SETUP_SPECS, import_setup_rows, validate_setup_workbook


EXPECTED_HIERARCHY = {
    "Science and Technology": {"Computer Science and Engineering", "Business Information Technology", "Data Science"},
    "Languages": {"English", "Chinese Language"},
    "Business": {"Business Administration", "Marketing"},
}


class DemoDataTests(TestCase):
    def test_generator_creates_only_new_physical_tree(self):
        original = Path.cwd()
        with TemporaryDirectory() as temporary:
            try:
                os.chdir(temporary)
                call_command("generate_demo_files", verbosity=0)
                root = Path("demo_data")
                expected = {
                    *(Path("Good Data") / f"{kind}.xlsx" for kind in SETUP_SPECS),
                    *(Path("Good Data") / f"{kind}.xlsx" for kind in ("enrollment", "performance", "attendance")),
                    *(Path("Messy Data") / f"{kind}_messy.xlsx" for kind in SETUP_SPECS),
                    *(Path("Messy Data") / f"{kind}_messy.xlsx" for kind in ("enrollment", "performance", "attendance")),
                    Path("README.md"),
                }
                self.assertEqual({path.relative_to(root) for path in root.rglob("*") if path.is_file()}, expected)
                for path in root.rglob("*.xlsx"):
                    self.assertGreater(load_workbook(path, read_only=True).active.max_row, 1)
                self.assertFalse((root / "demo_master_data.xlsx").exists())
                self.assertIn("department.xlsx", (root / "README.md").read_text())
                with self.assertRaises(CommandError):
                    call_command("generate_demo_files", verbosity=0)
            finally:
                os.chdir(original)

    def test_good_rows_import_and_messy_batches_show_valid_and_rejected_rows(self):
        good = demo_workbooks()
        messy = messy_demo_workbooks()
        order = ("department", "major", "subject", "teacher", "room", "enrollment", "performance", "attendance")
        self.assertEqual(set(good), set(order))
        self.assertEqual(set(messy), set(order))
        for kind in order:
            payload = BytesIO(good[kind][1])
            summary, valid, issues = (validate_setup_workbook(payload, kind) if kind in SETUP_SPECS
                                      else validate_workbook(payload, kind))
            self.assertEqual(summary["rejected_rows"], 0, (kind, issues))
            self.assertEqual(summary["valid_rows"], summary["total_rows"])
            if kind in SETUP_SPECS:
                import_setup_rows(kind, valid)
            else:
                import_rows(SimpleNamespace(dataset_type=kind, valid_rows=valid))
        self.assertEqual(Student.objects.count(), 227)
        original_name = Student.objects.get(student_id="SHOW0001").student_name
        for kind in order:
            payload = BytesIO(messy[kind][1])
            summary, valid, issues = (validate_setup_workbook(payload, kind) if kind in SETUP_SPECS
                                      else validate_workbook(payload, kind))
            self.assertGreater(summary["valid_rows"], 0, (kind, issues))
            self.assertGreater(summary["rejected_rows"], 0, kind)
            self.assertEqual(len(issues), summary["rejected_rows"])
            self.assertGreater(summary["duplicate_rows"], 0, kind)
            if kind == "enrollment":
                self.assertTrue(any("Existing student_id conflicts" in error for issue in issues for error in issue["errors"]))
            elif kind == "performance":
                self.assertTrue(any("score conflicts" in error for issue in issues for error in issue["errors"]))
            elif kind == "attendance":
                self.assertTrue(any("attendance conflicts" in error for issue in issues for error in issue["errors"]))
            if kind in SETUP_SPECS:
                import_setup_rows(kind, valid)
            else:
                import_rows(SimpleNamespace(dataset_type=kind, valid_rows=valid))
        self.assertEqual(Student.objects.filter(student_id__in=["SHOW9001", "SHOW9002", "SHOW9003"]).count(), 3)
        self.assertEqual(Student.objects.get(student_id="SHOW0001").student_name, original_name)

    def test_workbooks_use_requested_hierarchy_and_consistent_fictional_names(self):
        departments = {code: name for code, name in MASTER_ROWS["Departments"][1:]}
        hierarchy = {name: set() for name in departments.values()}
        for _, major_name, department_code in MASTER_ROWS["Majors"][1:]:
            hierarchy[departments[department_code]].add(major_name)
        self.assertEqual(hierarchy, EXPECTED_HIERARCHY)

        files = demo_workbooks()
        enrollment = list(load_workbook(BytesIO(files["enrollment"][1]), read_only=True).active.values)
        performance = list(load_workbook(BytesIO(files["performance"][1]), read_only=True).active.values)
        attendance = list(load_workbook(BytesIO(files["attendance"][1]), read_only=True).active.values)
        names_by_id = {student_id: name for student_id, name, *_ in enrollment[1:]}
        self.assertEqual(len(names_by_id), len(enrollment) - 1)
        self.assertEqual(len(names_by_id), 227)
        self.assertGreater(len(set(names_by_id.values())), 200)
        self.assertTrue(all(len(name.split()) == 2 for name in names_by_id.values()))
        self.assertFalse(any("Student" in name or "Test" in name or "Doe" in name or "Smith" in name for name in names_by_id.values()))
        self.assertEqual(set(row[0] for row in performance[1:]), set(names_by_id))
        self.assertEqual(set(row[0] for row in attendance[1:]), set(names_by_id))
        self.assertEqual({tuple(row[:2]) for row in performance[1:]}, {tuple(row[:2]) for row in attendance[1:]})
        self.assertEqual([row[1] for row in enrollment[1:]], cambodian_student_names(227))
        self.assertEqual(
            list(load_workbook(BytesIO(demo_workbooks()["enrollment"][1]), read_only=True).active.values),
            enrollment,
        )

    def test_seed_demo_keeps_names_stable_across_reruns_and_files(self):
        original = Path.cwd()
        with TemporaryDirectory() as temporary:
            try:
                os.chdir(temporary)
                call_command("seed_demo", students=35, seed=11, verbosity=0)
                first_names = dict(Student.objects.values_list("student_id", "student_name"))
                call_command("seed_demo", students=35, seed=99, verbosity=0)
                self.assertEqual(dict(Student.objects.values_list("student_id", "student_name")), first_names)
                enrollment = list(load_workbook(Path("demo_data/synthetic_enrollment.xlsx"), read_only=True).active.values)
                performance = list(load_workbook(Path("demo_data/synthetic_performance.xlsx"), read_only=True).active.values)
                attendance = list(load_workbook(Path("demo_data/synthetic_attendance.xlsx"), read_only=True).active.values)
                self.assertEqual({row[0]: row[1] for row in enrollment[1:]}, first_names)
                self.assertTrue(set(row[0] for row in performance[1:]) <= set(first_names))
                self.assertTrue(set(row[0] for row in attendance[1:]) <= set(first_names))
                self.assertEqual({department.name: set(department.majors.values_list("name", flat=True))
                                  for department in Department.objects.all()}, EXPECTED_HIERARCHY)
                self.assertEqual(Major.objects.count(), 7)
            finally:
                os.chdir(original)
