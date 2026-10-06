"""Regression coverage for individual setup uploads and native Excel reports."""

from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

from .models import Attendance, Department, Major, Performance, Room, Student, Subject, Teacher, User
from .services import attendance_status, grade_for
from .setup_import import SETUP_SPECS


def excel_file(headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return SimpleUploadedFile("setup.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class SetupImportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="setup@example.com", full_name="Setup Analyst", password="SafePass123!")
        self.client.force_login(self.user)

    def preview(self, kind, rows):
        response = self.client.post(reverse("setup_import", args=[kind]), {"file": excel_file(SETUP_SPECS[kind], rows)})
        self.assertEqual(response.status_code, 302)
        review = self.client.get(response.url)
        self.assertEqual(review.status_code, 200)
        return response.url, review.context["batch"]

    def confirm(self, url):
        response = self.client.post(url + "confirm/", follow=True)
        self.assertEqual(response.status_code, 200)
        return response.context["batch"]

    def test_all_five_templates_and_preview_confirm_flow(self):
        examples = {
            "department": [["CSE", "Computer Science"]],
            "major": [["SE", "Software Engineering", "CSE"]],
            "subject": [["SE101", "Programming", "SE"]],
            "teacher": [["Dara Sok", "dara@example.com", "CSE", "Morning,Afternoon", "SE101", "true"]],
            "room": [["A101", 25, "true"]],
        }
        for kind in examples:
            template = self.client.get(reverse("setup_template", args=[kind]))
            self.assertEqual(template.status_code, 200)
            self.assertIn("application/vnd.openxmlformats", template["Content-Type"])
            self.assertEqual(list(load_workbook(BytesIO(template.content)).active.values)[0], tuple(SETUP_SPECS[kind]))
            listing = self.client.get(reverse(f"{kind}_list"))
            self.assertContains(listing, "Import Excel")
            self.assertContains(listing, "Download Template")
            url, batch = self.preview(kind, examples[kind])
            self.assertEqual((batch.summary["valid_rows"], batch.summary["rejected_rows"]), (1, 0))
            self.assertEqual({"department": Department, "major": Major, "subject": Subject,
                              "teacher": Teacher, "room": Room}[kind].objects.count(), 0)
            self.assertIsNotNone(self.confirm(url).imported_at)
        teacher = Teacher.objects.get(email="dara@example.com")
        self.assertEqual(teacher.available_shifts, ["Morning", "Afternoon"])
        self.assertEqual(list(teacher.subjects_can_teach.values_list("code", flat=True)), ["SE101"])
        self.assertEqual(Room.objects.get(code="A101").capacity, 25)
        self.assertNotContains(self.client.get(reverse("dashboard")), "Import Master Data")
        self.assertEqual(self.client.get("/upload/master/").status_code, 404)

    def test_downloaded_demo_templates_import_in_dependency_order(self):
        for kind in ("department", "major", "subject", "teacher", "room"):
            payload = self.client.get(reverse("setup_template", args=[kind])).content
            response = self.client.post(reverse("setup_import", args=[kind]),
                {"file": SimpleUploadedFile(f"{kind}_template.xlsx", payload)})
            self.assertEqual(response.status_code, 302)
            batch = self.client.get(response.url).context["batch"]
            self.assertEqual(batch.summary["rejected_rows"], 0, (kind, batch.issues))
            self.assertIsNotNone(self.confirm(response.url).imported_at)
        self.assertEqual((Department.objects.count(), Major.objects.count(), Subject.objects.count(),
                          Teacher.objects.count(), Room.objects.count()), (3, 7, 11, 6, 4))

    def test_setup_uploads_and_templates_require_login(self):
        self.client.logout()
        for kind in ("department", "major", "subject", "teacher", "room"):
            for route in ("setup_import", "setup_template"):
                url = reverse(route, args=[kind])
                self.assertRedirects(self.client.get(url), f"{reverse('login')}?next={url}")

    def test_rejections_are_reviewable_and_existing_records_unchanged(self):
        Department.objects.create(code="CSE", name="Original")
        url, batch = self.preview("department", [["CSE", "Overwrite Attempt"], ["BUS", "Business"],
                                                  ["BUS", "Duplicate"], ["", "Missing Code"]])
        self.assertEqual((batch.summary["total_rows"], batch.summary["valid_rows"],
                          batch.summary["duplicate_rows"], batch.summary["missing_value_rows"]), (4, 1, 2, 1))
        self.assertEqual({issue["row"] for issue in batch.issues}, {2, 4, 5})
        self.confirm(url)
        self.assertEqual(Department.objects.get(code="CSE").name, "Original")
        self.assertEqual(Department.objects.get(code="BUS").name, "Business")
        bad = self.client.post(reverse("setup_import", args=["major"]),
                               {"file": excel_file(["code", "name", "department_code"], [["SE", "Software Engineering", "UNKNOWN"]])})
        self.assertEqual(self.client.get(bad.url).context["batch"].summary["invalid_references"], 1)
        bad_subject = self.client.post(reverse("setup_import", args=["subject"]),
            {"file": excel_file(["code", "name", "major_code"], [["S1", "Subject", "UNKNOWN"]])})
        self.assertEqual(self.client.get(bad_subject.url).context["batch"].summary["invalid_references"], 1)
        corrupt = self.client.post(reverse("setup_import", args=["room"]),
            {"file": SimpleUploadedFile("bad.xlsx", b"not-a-workbook")})
        self.assertContains(corrupt, "Excel file could not be read")
        missing_column = self.client.post(reverse("setup_import", args=["room"]),
            {"file": excel_file(["code"], [["A101"]])})
        self.assertContains(missing_column, "Missing required columns")

    def test_teacher_subject_and_shift_rules_and_room_validation(self):
        cse = Department.objects.create(code="CSE", name="Computer Science")
        bus = Department.objects.create(code="BUS", name="Business")
        se = Major.objects.create(code="SE", name="Software Engineering", department=cse)
        bba = Major.objects.create(code="BBA", name="Business", department=bus)
        for code in ("S1", "S2", "S3", "S4"):
            Subject.objects.create(code=code, name=code, major=se)
        Subject.objects.create(code="B1", name="Business", major=bba)
        rows = [
            ["One", "one@example.com", "CSE", "Morning", "S1,S2,S3,S4", "true"],
            ["Two", "two@example.com", "CSE", "Morning", "S1,B1", "true"],
            ["Three", "three@example.com", "CSE", "Morning,Weekend", "S1", "true"],
            ["Four", "four@example.com", "CSE", "Evening", "S1", "maybe"],
            ["Five", "five@example.com", "CSE", "Morning,Afternoon", "S1,S2", "false"],
        ]
        url, batch = self.preview("teacher", rows)
        self.assertEqual((batch.summary["valid_rows"], batch.summary["rejected_rows"]), (1, 4))
        self.assertTrue(any("3 subjects" in str(issue["errors"]) for issue in batch.issues))
        self.assertTrue(any("selected department" in str(issue["errors"]) or "valid choice" in str(issue["errors"]) for issue in batch.issues))
        self.confirm(url)
        self.assertFalse(Teacher.objects.get(name="Five").active)
        _, rooms = self.preview("room", [["R1", 0, "true"], ["R2", 25, "maybe"], ["R3", 20, "false"]])
        self.assertEqual((rooms.summary["valid_rows"], rooms.summary["invalid_numeric_rows"]), (1, 1))


class ModuleReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="report@example.com", full_name="Report Analyst", password="SafePass123!")
        self.client.force_login(self.user)

    def report(self, kind, query=""):
        response = self.client.get(reverse("export_analytics", args=[kind]) + query)
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats", response["Content-Type"])
        self.assertIn("202", response["Content-Disposition"])
        return load_workbook(BytesIO(response.content))

    def test_empty_reports_have_no_broken_charts(self):
        sheets = {
            "enrollment": ["Overview", "Enrollment Data", "Major & Shift Analysis", "Capacity Analysis", "Insights", "_ChartData"],
            "performance": ["Overview", "Student Performance", "Subject Analysis", "Grade Distribution", "Insights", "_ChartData"],
            "attendance": ["Overview", "Attendance Records", "Subject Attendance", "Risk Analysis", "Attendance vs Performance", "Insights", "_ChartData"],
        }
        for kind in sheets:
            book = self.report(kind)
            self.assertEqual(book.sheetnames, sheets[kind])
            self.assertEqual(book["_ChartData"].sheet_state, "hidden")
            self.assertEqual(book["Overview"]._charts, [])

    def test_partial_data_omits_unsupported_charts(self):
        department = Department.objects.create(code="CSE", name="Computer Science")
        major = Major.objects.create(code="SE", name="Software Engineering", department=department)
        subject = Subject.objects.create(code="SE101", name="Programming", major=major)
        student = Student.objects.create(student_id="ST001", student_name="Student One",
                                         department=department, major=major, shift="Morning")
        self.assertEqual(len(self.report("enrollment")["Overview"]._charts), 6)
        self.assertEqual(self.report("performance")["Overview"]._charts, [])
        self.assertEqual(self.report("attendance")["Overview"]._charts, [])
        Performance.objects.create(student=student, subject=subject, score=78,
                                   letter_grade="B", grade_point=3, passed=True)
        self.assertEqual(len(self.report("performance")["Overview"]._charts), 5)
        self.assertEqual(self.report("attendance")["Overview"]._charts, [])
        Performance.objects.all().delete()
        Attendance.objects.create(student=student, subject=subject, total_sessions=20,
                                  absent_count=2, attendance_percentage=90, status="Good")
        attendance = self.report("attendance")
        self.assertEqual(len(attendance["Overview"]._charts), 4)
        self.assertFalse(any(type(chart).__name__ == "ScatterChart" for chart in attendance["Overview"]._charts))

    def test_full_and_filtered_reports_have_real_numeric_chart_sources(self):
        cse = Department.objects.create(code="CSE", name="Computer Science")
        bus = Department.objects.create(code="BUS", name="Business")
        se = Major.objects.create(code="SE", name="Software Engineering", department=cse)
        bba = Major.objects.create(code="BBA", name="Business", department=bus)
        database = Subject.objects.create(code="DB101", name="Database", major=se)
        business = Subject.objects.create(code="B101", name="Business", major=bba)
        Teacher.objects.create(name="Dara", department=cse, available_shifts=["Morning"], active=True)
        Room.objects.create(code="A101", capacity=25, active=True)
        students = [Student.objects.create(student_id=f"ST{i:03}", student_name=f"Student {i}",
                   department=cse if i < 3 else bus, major=se if i < 3 else bba,
                   shift="Morning" if i < 3 else "Evening") for i in range(1, 4)]
        for student, subject, score, absent in [(students[0], database, 100, 0),
                                                 (students[1], database, 40, 9),
                                                 (students[2], business, 80, 2)]:
            grade, point = grade_for(score)
            Performance.objects.create(student=student, subject=subject, score=score,
                                       letter_grade=grade, grade_point=point, passed=score >= 50)
            percentage = (20 - absent) / 20 * 100
            Attendance.objects.create(student=student, subject=subject, total_sessions=20,
                                      absent_count=absent, attendance_percentage=percentage,
                                      status=attendance_status(percentage))
        expected = {"enrollment": 6, "performance": 5, "attendance": 5}
        for kind, chart_count in expected.items():
            book = self.report(kind)
            self.assertEqual(len(book["Overview"]._charts), chart_count)
            for chart in book["Overview"]._charts:
                self.assertFalse(chart.visible_cells_only)
                self.assertTrue(chart.series)
                for series in chart.series:
                    refs = [ref for ref in (getattr(series, "val", None), getattr(series, "xVal", None), getattr(series, "yVal", None)) if ref]
                    self.assertTrue(refs)
                    for ref in refs:
                        formula = ref.numRef.f if ref.numRef else None
                        self.assertIn("_ChartData", formula)
                        self.assertNotIn("#REF!", formula)
                        bounds = range_boundaries(formula.split("!")[-1])
                        self.assertLessEqual(bounds[2], book["_ChartData"].max_column)
                        self.assertLessEqual(bounds[3], book["_ChartData"].max_row)
                        self.assertIsInstance(book["_ChartData"].cell(bounds[1], bounds[0]).value, (int, float))
            self.assertEqual(book["Overview"]["C6"].data_type, "n")
        filtered = f"?department={cse.pk}&major={se.pk}&shift=Morning"
        enrollment = self.report("enrollment", filtered)
        self.assertEqual(enrollment["Enrollment Data"].max_row, 3)
        self.assertIn("Department: CSE", enrollment["Overview"]["C4"].value)
        self.assertEqual(enrollment["Overview"]["C6"].value, 2)
        self.assertEqual(enrollment["Capacity Analysis"]["G2"].value, 1)
        self.assertEqual(enrollment["Capacity Analysis"]["J2"].value, 1)
        performance = self.report("performance", filtered)
        self.assertEqual(performance["Student Performance"].max_row, 3)
        self.assertEqual(performance["Student Performance"]["G2"].data_type, "n")
        attendance = self.report("attendance", filtered)
        self.assertEqual(attendance["Attendance Records"].max_row, 3)
        self.assertEqual(attendance["Attendance Records"]["I2"].data_type, "n")
        scatter = attendance["Overview"]._charts[-1]
        self.assertEqual(type(scatter).__name__, "ScatterChart")
        self.assertIn("_ChartData", scatter.series[0].xVal.numRef.f)
        self.assertIn("_ChartData", scatter.series[0].yVal.numRef.f)
        self.assertEqual((Student.objects.count(), Performance.objects.count(), Attendance.objects.count()), (3, 3, 3))
