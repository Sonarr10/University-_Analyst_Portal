from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from .models import Attendance, Department, Major, Performance, Room, Shift, Student, Subject, Teacher, User
from .capacity import MAX_STUDENTS_PER_SECTION, capacity_analysis, required_sections, section_occupancy
from .demo_workbooks import demo_workbooks
from .forms import TeacherForm
from .master_import import validate_master_workbook
from .services import capacity_rows, dashboard_context, validate_workbook


def workbook(headers, rows):
    book = Workbook(); sheet = book.active; sheet.append(headers)
    for row in rows: sheet.append(row)
    data = BytesIO(); book.save(data); data.seek(0); data.name = "test.xlsx"; return data


class PortalTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="analyst@example.com", full_name="Demo Analyst", password="SafePass123!")
        self.department = Department.objects.create(code="CSE", name="Computer Science")
        self.major = Major.objects.create(code="SE", name="Software Engineering", department=self.department)
        self.subject = Subject.objects.create(code="DB101", name="Database Systems", major=self.major)

    def test_internal_pages_require_login_and_registration_hashes_password(self):
        self.assertRedirects(self.client.get(reverse("dashboard")), f"{reverse('login')}?next=/")
        response = self.client.post(reverse("register"), {"full_name": "New Analyst", "email": "new@example.com", "password1": "AnotherSafe123!", "password2": "AnotherSafe123!"})
        self.assertRedirects(response, reverse("dashboard"))
        self.assertTrue(User.objects.get(email="new@example.com").check_password("AnotherSafe123!"))

    def test_crud_and_empty_dashboard(self):
        self.client.login(username="analyst@example.com", password="SafePass123!")
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, "Build your analytics workspace")
        response = self.client.post(reverse("room_create"), {"code": "R01", "capacity": 30, "active": "on"})
        self.assertRedirects(response, reverse("room_list")); self.assertTrue(Room.objects.filter(code="R01").exists())

    def test_teacher_shift_checkboxes_use_group_layout(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("teacher_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<div id="id_available_shifts" class="checkbox-choices">')
        self.assertNotContains(response, '<div id="id_available_shifts" class="form-check-input">')
        self.assertContains(response, '<fieldset>')
        self.assertNotContains(response, "Max sections")
        self.assertNotContains(response, "Assigned sections")
        self.assertContains(response, "Selected: 0 / 3")

    def test_enrollment_cleaning_and_confirmed_import(self):
        self.client.login(username="analyst@example.com", password="SafePass123!")
        file = workbook(["student_id", "student_name", "department_code", "major_code", "shift"], [
            ["ST001", "Valid Student", "CSE", "SE", "Morning"],
            ["ST001", "Duplicate", "CSE", "SE", "Morning"],
            ["ST003", "Wrong Shift", "CSE", "SE", "Weekend"],
        ])
        response = self.client.post(reverse("upload_data"), {"dataset_type": "enrollment", "file": file})
        self.assertEqual(response.status_code, 302)
        batch_url = response.url
        response = self.client.get(batch_url); self.assertContains(response, "2 rejected")
        response = self.client.post(batch_url + "confirm/")
        self.assertEqual(Student.objects.count(), 1)

    def test_performance_and_attendance_validation_import(self):
        student = Student.objects.create(student_id="ST001", student_name="Student", department=self.department, major=self.major, shift=Shift.MORNING)
        p_summary, p_rows, p_issues = validate_workbook(workbook(["student_id", "subject_code", "score"], [["ST001", "DB101", 86], ["BAD", "DB101", 150]]), "performance")
        self.assertEqual(p_summary["valid_rows"], 1); self.assertEqual(len(p_issues), 1)
        a_summary, a_rows, a_issues = validate_workbook(workbook(["student_id", "subject_code", "total_sessions", "absent_count"], [["ST001", "DB101", 20, 4], ["ST001", "DB101", 10, 11]]), "attendance")
        self.assertEqual(a_summary["valid_rows"], 1); self.assertEqual(len(a_issues), 1)

    def test_performance_and_attendance_upload_flow(self):
        Student.objects.create(student_id="ST001", student_name="Student", department=self.department, major=self.major, shift=Shift.MORNING)
        self.client.login(username="analyst@example.com", password="SafePass123!")
        cases = [
            ("performance", ["student_id", "subject_code", "score"], [["ST001", "DB101", 86]], Performance),
            ("attendance", ["student_id", "subject_code", "total_sessions", "absent_count"], [["ST001", "DB101", 20, 4]], Attendance),
        ]
        for dataset_type, headers, rows, model in cases:
            response = self.client.post(reverse("upload_data"), {"dataset_type": dataset_type, "file": workbook(headers, rows)})
            self.assertEqual(response.status_code, 302)
            self.client.post(response.url + "confirm/")
            self.assertEqual(model.objects.count(), 1)
        self.assertEqual(Performance.objects.get().letter_grade, "A")
        self.assertEqual(Attendance.objects.get().status, "Warning")

    def test_chart_assets_only_render_for_available_data(self):
        self.client.login(username="analyst@example.com", password="SafePass123!")
        response = self.client.get(reverse("dashboard"))
        self.assertNotContains(response, 'id="departmentChart"')
        Student.objects.create(student_id="ST001", student_name="Student", department=self.department, major=self.major, shift=Shift.MORNING)
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, 'id="departmentChart"')
        self.assertContains(response, "portalCharts.dashboard()")

    def test_capacity_calculation_and_dashboards(self):
        Room.objects.create(code="R01", capacity=30, active=True)
        Teacher.objects.create(name="Teacher", department=self.department, available_shifts=[Shift.MORNING], max_sections=2)
        for i in range(26):
            Student.objects.create(student_id=f"ST{i:03}", student_name=f"Student {i}", department=self.department, major=self.major, shift=Shift.MORNING)
        row = capacity_rows()[0]
        self.assertEqual(row["required"], 2)
        self.assertEqual(row["teacher_capacity"], 1)
        self.assertEqual(row["teacher_shortage"], 1)
        self.assertEqual(row["room_shortage"], 1)
        self.assertEqual(row["status"], "Over Capacity")
        self.client.login(username="analyst@example.com", password="SafePass123!")
        for url in ["dashboard", "enrollment_dashboard", "performance_dashboard", "attendance_dashboard"]:
            self.assertEqual(self.client.get(reverse(url)).status_code, 200)


class DashboardPresentationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="cards@example.com", full_name="Card Analyst", password="SafePass123!")
        self.client.force_login(self.user)
        self.department = Department.objects.create(code="CSE", name="Computer Science")
        self.major = Major.objects.create(code="SE", name="Software Engineering", department=self.department)
        self.teacher = Teacher.objects.create(name="Teacher", department=self.department, available_shifts=[Shift.MORNING])
        Room.objects.bulk_create([Room(code="R1", capacity=25), Room(code="R2", capacity=25)])

    def enroll(self, size, shift=Shift.MORNING, major=None):
        Student.objects.bulk_create([
            Student(student_id=f"{shift}-{size}-{index}", student_name="Student", department=self.department, major=major or self.major, shift=shift)
            for index in range(size)
        ])

    def test_student_card_stays_neutral_at_all_capacity_levels(self):
        for size in (18, 20, 22, 23, 24, 25, 26, 51):
            with self.subTest(students=size):
                Student.objects.all().delete()
                self.enroll(size)
                response = self.client.get(reverse("dashboard"))
                self.assertEqual(response.status_code, 200)
                student_card = response.content.decode().split('<div class="kpi-card">', 1)[1].split('<div class="kpi-card">', 1)[0]
                self.assertIn(f"<strong>{size}</strong>", student_card)
                for warning in ("Full", "Near Full", "Getting Full", "remaining", "kpi-status"):
                    self.assertNotIn(warning, student_card)
                self.assertNotIn("full_sections", response.context)
                self.assertNotIn("near_full_sections", response.context)

        Student.objects.all().delete()
        self.enroll(25)
        second_major = Major.objects.create(code="CYB", name="Cybersecurity", department=self.department)
        self.enroll(23, Shift.AFTERNOON, second_major)
        response = self.client.get(reverse("dashboard"))
        student_card = response.content.decode().split('<div class="kpi-card">', 1)[1].split('<div class="kpi-card">', 1)[0]
        self.assertNotIn("kpi-status", student_card)

    def test_teacher_badge_and_compact_insights_states(self):
        self.enroll(18)
        healthy = self.client.get(reverse("dashboard"))
        self.assertNotContains(healthy, "⚠ Shortage")
        self.assertEqual(healthy.context["dashboard_insights"], [])
        self.assertContains(healthy, "No critical issues detected.")

        self.teacher.active = False
        self.teacher.save(update_fields=["active"])
        shortage = self.client.get(reverse("dashboard"))
        self.assertContains(shortage, '<span class="kpi-status kpi-status-warning">⚠ Shortage</span>')
        self.assertEqual(len(shortage.context["dashboard_insights"]), 1)
        self.assertContains(shortage, "Teacher shortage in CSE Morning")
        self.assertContains(shortage, 'href="/insights/"')
        details = self.client.get(reverse("insights"))
        self.assertContains(details, "1 additional teacher(s) required")
        self.assertContains(details, "Teacher Capacity")

        Student.objects.all().delete()
        self.enroll(51)
        critical = self.client.get(reverse("dashboard"))
        self.assertContains(critical, '<span class="kpi-status kpi-status-danger">⚠ Shortage</span>')

    def test_insights_page_requires_login(self):
        self.client.logout()
        self.assertRedirects(self.client.get(reverse("insights")), f"{reverse('login')}?next={reverse('insights')}")


class TeacherDemoFlowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="demo@example.com", full_name="Demo Analyst", password="SafePass123!")
        self.client.force_login(self.user)
        self.files = demo_workbooks()

    def upload_and_confirm(self, kind):
        filename, payload = self.files[kind]
        endpoint = "import_master_data" if kind == "master" else "upload_data"
        fields = {"file": SimpleUploadedFile(filename, payload, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        if kind != "master":
            fields["dataset_type"] = kind
        response = self.client.post(reverse(endpoint), fields)
        self.assertEqual(response.status_code, 302, response.content.decode()[:500])
        review = self.client.get(response.url)
        self.assertEqual(review.status_code, 200)
        confirmed = self.client.post(response.url + "confirm/", follow=True)
        self.assertEqual(confirmed.status_code, 200)
        self.assertTrue(confirmed.context["batch"].imported_at)
        return review

    def test_complete_teacher_walkthrough(self):
        empty = self.client.get(reverse("dashboard"))
        self.assertContains(empty, "Build your analytics workspace")
        self.assertNotContains(empty, 'id="departmentChart"')
        self.assertContains(empty, "No Data")

        self.assertRedirects(self.client.post(reverse("department_create"), {"code": "CSE", "name": "Computer Science and Engineering"}), reverse("department_list"))
        self.assertRedirects(self.client.post(reverse("major_create"), {"code": "SE", "name": "Software Engineering", "department": Department.objects.get(code="CSE").pk}), reverse("major_list"))
        setup_progress = self.client.get(reverse("dashboard"))
        self.assertTrue(setup_progress.context["onboarding_steps"][0]["done"])
        self.assertTrue(setup_progress.context["onboarding_steps"][1]["done"])
        self.assertEqual(self.client.get(reverse("upload_data") + "?dataset_type=performance").context["form"].initial["dataset_type"], "performance")
        master_review = self.upload_and_confirm("master")
        self.assertContains(master_review, "Departments")
        self.assertEqual((Department.objects.count(), Major.objects.count(), Subject.objects.count(), Teacher.objects.count(), Room.objects.count()), (3, 4, 7, 6, 4))

        self.upload_and_confirm("master")
        self.assertEqual((Department.objects.count(), Major.objects.count(), Subject.objects.count(), Teacher.objects.count(), Room.objects.count()), (3, 4, 7, 6, 4))
        self.upload_and_confirm("enrollment")
        self.assertEqual(Student.objects.count(), 185)
        self.upload_and_confirm("performance")
        self.upload_and_confirm("attendance")
        self.assertEqual(Performance.objects.count(), 397)
        self.assertEqual(Attendance.objects.count(), 397)

        analysis = capacity_analysis()
        se_morning = next(row for row in analysis["groups"] if row["major__name"] == "Software Engineering" and row["shift"] == "Morning")
        self.assertEqual((se_morning["students"], se_morning["required"], se_morning["teacher_capacity"], se_morning["status"]), (70, 3, 2, "Over Capacity"))
        cse_morning = next(pool for pool in analysis["teacher_pools"] if pool["department"] == "CSE" and pool["shift"] == "Morning")
        self.assertEqual((cse_morning["required"], cse_morning["capacity"], cse_morning["shortage"]), (5, 2, 3))
        bba_afternoon = next(row for row in analysis["groups"] if row["major__name"] == "Business Administration" and row["shift"] == "Afternoon")
        self.assertEqual(bba_afternoon["status"], "Near Capacity")
        se_afternoon = next(row for row in analysis["groups"] if row["major__name"] == "Software Engineering" and row["shift"] == "Afternoon")
        self.assertEqual(se_afternoon["status"], "Teacher Shortage")
        bba_morning = next(row for row in analysis["groups"] if row["major__name"] == "Business Administration" and row["shift"] == "Morning")
        self.assertEqual(bba_morning["status"], "Room Capacity Limit")
        english_evening = next(row for row in analysis["groups"] if row["major__name"] == "English Language" and row["shift"] == "Evening")
        self.assertEqual(english_evening["status"], "Available")
        self.assertTrue(any(section["status"] == "Full" for section in analysis["occupancy"]))
        self.assertTrue(any(section["status"] == "Near Full" for section in analysis["occupancy"]))
        self.assertTrue(any(section["status"] == "Getting Full" for section in analysis["occupancy"]))
        self.assertTrue(any(section["status"] == "Normal" for section in analysis["occupancy"]))

        context = dashboard_context()
        insight_text = " ".join(item["text"] for item in context["insights"])
        self.assertIn("Database Systems has the lowest average score", insight_text)
        self.assertIn("12 students have at least one Critical attendance record", insight_text)
        self.assertIn("Matched low-attendance results average", insight_text)
        self.assertIn("CSE Morning requires 5 teachers", insight_text)
        self.assertIn("Morning requires 6 simultaneous shared rooms", insight_text)
        self.assertNotIn("create another room", insight_text.lower())
        self.assertEqual(context["kpis"]["students"], 185)
        self.assertEqual(len(context["availability_alerts"]), 3)
        self.assertEqual(context["availability_alerts"][0]["status"], "Over Capacity")
        previews = context["dashboard_insights"]
        self.assertGreater(len(context["insights"]), 5)
        self.assertEqual(len(previews), 5)
        self.assertTrue(all(len(item["summary"]) < 80 for item in previews))
        self.assertEqual([{"danger": 0, "warning": 1, "info": 2}[item["level"]] for item in previews], sorted({"danger": 0, "warning": 1, "info": 2}[item["level"]] for item in previews))
        self.assertTrue(any("Teacher shortage in CSE Morning" == item["summary"] for item in previews))
        self.assertEqual(context["insights"][0]["level"], "danger")
        self.assertEqual([{"danger": 0, "warning": 1, "info": 2}[item["level"]] for item in context["insights"]], sorted({"danger": 0, "warning": 1, "info": 2}[item["level"]] for item in context["insights"]))
        self.assertContains(self.client.get(reverse("insights")), "5 sections across its majors")
        for name in ("dashboard", "enrollment_dashboard", "performance_dashboard", "attendance_dashboard"):
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, "portalCharts.")

    def test_master_workbook_rejects_bad_relationships_and_values(self):
        filename, payload = self.files["master"]
        book = load_workbook(BytesIO(payload))
        book["Majors"]["C2"] = "UNKNOWN"
        book["Teachers"]["D2"] = "Morning,Weekend"
        output = BytesIO(); book.save(output); output.seek(0)
        summary, valid, issues = validate_master_workbook(output)
        self.assertGreaterEqual(summary["rejected_rows"], 2)
        self.assertTrue(any("Unknown department_code" in error for issue in issues for error in issue["errors"]))
        self.assertTrue(any("Available shifts" in error for issue in issues for error in issue["errors"]))
        del book["Rooms"]
        missing = BytesIO(); book.save(missing); missing.seek(0)
        with self.assertRaisesRegex(ValueError, "Missing required sheets: Rooms"):
            validate_master_workbook(missing)

    def test_demo_downloads_are_real_workbooks(self):
        for kind, (filename, _) in self.files.items():
            response = self.client.get(reverse("download_demo", args=[kind]))
            self.assertEqual(response.status_code, 200)
            self.assertIn(filename, response["Content-Disposition"])
            workbook = load_workbook(BytesIO(response.content), read_only=True)
            self.assertTrue(workbook.sheetnames)


class TeacherCapacityRuleTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="rules@example.com", full_name="Rules Analyst", password="SafePass123!")
        self.client.force_login(self.user)
        self.cse = Department.objects.create(code="CSE", name="Computer Science")
        self.bus = Department.objects.create(code="BUS", name="Business")
        self.se = Major.objects.create(code="SE", name="Software Engineering", department=self.cse)
        self.bba = Major.objects.create(code="BBA", name="Business Administration", department=self.bus)
        self.cse_subjects = [Subject.objects.create(code=f"SE{i}", name=f"SE Subject {i}", major=self.se) for i in range(4)]
        self.bus_subject = Subject.objects.create(code="BUS1", name="Business Subject", major=self.bba)

    def teacher_payload(self, department, shifts, subjects, name="Teacher One"):
        return {
            "name": name, "email": "teacher@example.com", "department": department.pk,
            "available_shifts": shifts, "subjects_can_teach": [str(subject.pk) for subject in subjects],
            "active": "on",
        }

    def test_create_edit_department_filter_and_subject_limit(self):
        response = self.client.get(reverse("teacher_create"))
        self.assertEqual(list(response.context["form"].fields["subjects_can_teach"].queryset), [])
        self.assertContains(response, 'id="teacher-subject-data"')
        self.assertContains(response, "js/teacher_form.js")

        payload = self.teacher_payload(self.cse, [Shift.MORNING, Shift.AFTERNOON], self.cse_subjects[:2])
        response = self.client.post(reverse("teacher_create"), payload)
        self.assertRedirects(response, reverse("teacher_list"))
        teacher = Teacher.objects.get(name="Teacher One")
        self.assertEqual(teacher.available_shifts, [Shift.MORNING, Shift.AFTERNOON])
        self.assertEqual(teacher.subjects_can_teach.count(), 2)
        self.assertEqual(teacher.max_sections, 1)  # Legacy column is not an input or capacity source.

        edit_url = reverse("teacher_update", args=[teacher.pk])
        edit = self.client.get(edit_url)
        self.assertEqual(set(edit.context["form"].fields["subjects_can_teach"].queryset), set(self.cse_subjects))
        self.assertNotIn(self.bus_subject, edit.context["form"].fields["subjects_can_teach"].queryset)
        list_page = self.client.get(reverse("teacher_list"))
        self.assertContains(list_page, "Morning, Afternoon")
        self.assertContains(list_page, "SE Subject 0")
        self.assertNotContains(list_page, "Max sections")

        too_many = TeacherForm(data=self.teacher_payload(self.cse, [Shift.MORNING], self.cse_subjects))
        self.assertFalse(too_many.is_valid())
        self.assertIn("Select no more than 3 subjects", str(too_many.errors))
        too_many_post = self.client.post(reverse("teacher_create"), self.teacher_payload(self.cse, [Shift.MORNING], self.cse_subjects, "Too Many"))
        self.assertEqual(too_many_post.status_code, 200)
        self.assertFalse(Teacher.objects.filter(name="Too Many").exists())
        bad_create = self.client.post(reverse("teacher_create"), self.teacher_payload(self.cse, [Shift.MORNING], [self.bus_subject], "Bad Teacher"))
        self.assertEqual(bad_create.status_code, 200)
        self.assertFalse(Teacher.objects.filter(name="Bad Teacher").exists())

        wrong_department = self.client.post(edit_url, self.teacher_payload(self.bus, [Shift.EVENING], [self.cse_subjects[0]]))
        self.assertEqual(wrong_department.status_code, 200)
        teacher.refresh_from_db()
        self.assertEqual(teacher.department, self.cse)
        valid_edit = self.client.post(edit_url, self.teacher_payload(self.bus, [Shift.EVENING], [self.bus_subject]))
        self.assertRedirects(valid_edit, reverse("teacher_list"))
        teacher.refresh_from_db()
        self.assertEqual(teacher.department, self.bus)
        self.assertEqual(teacher.available_shifts, [Shift.EVENING])
        self.assertEqual(list(teacher.subjects_can_teach.all()), [self.bus_subject])

    def test_django_admin_also_hides_legacy_capacity_fields(self):
        superuser = User.objects.create_superuser(email="super@example.com", full_name="Super Analyst", password="SafePass123!")
        self.client.force_login(superuser)
        response = self.client.get("/admin/analytics/teacher/add/")
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Max sections")
        self.assertNotContains(response, "Assigned sections")

    def test_one_active_teacher_per_selected_shift(self):
        Room.objects.bulk_create([Room(code="R1", capacity=25), Room(code="R2", capacity=25)])
        Teacher.objects.create(name="Morning", department=self.cse, available_shifts=[Shift.MORNING], max_sections=99)
        Teacher.objects.create(name="Two shifts", department=self.cse, available_shifts=[Shift.MORNING, Shift.AFTERNOON], max_sections=99)
        Teacher.objects.create(name="All shifts", department=self.cse, available_shifts=Shift.values, max_sections=99)
        Teacher.objects.create(name="Inactive", department=self.cse, available_shifts=Shift.values, max_sections=99, active=False)
        Student.objects.bulk_create([
            Student(student_id=f"{shift[:1]}{index:03}", student_name="Student", department=self.cse, major=self.se, shift=shift)
            for shift in Shift.values for index in range(26)
        ])
        pools = {pool["shift"]: pool for pool in capacity_analysis()["teacher_pools"]}
        self.assertEqual([(pools[shift]["required"], pools[shift]["capacity"], pools[shift]["shortage"], pools[shift]["status"]) for shift in Shift.values], [
            (2, 3, 0, "Good"), (2, 2, 0, "At Capacity"), (2, 1, 1, "Teacher Shortage"),
        ])
        context = dashboard_context()
        self.assertEqual(len(context["teacher_shortage_pools"]), 1)
        self.assertEqual(context["teacher_shortage_pools"][0]["shift"], Shift.EVENING)
        response = self.client.get(reverse("enrollment_dashboard"))
        self.assertContains(response, "Need 1 more teacher")
        self.assertContains(response, "Room capacity")
        self.assertContains(response, 'id="roomChart"')
        self.assertContains(response, 'id="capacity-data"')
        self.assertContains(response, 'id="room-data"')
        dashboard = self.client.get(reverse("dashboard"))
        self.assertContains(dashboard, "⚠ Shortage")
        self.assertContains(dashboard, "Teacher shortage in CSE Evening")
        self.assertNotContains(dashboard, "CSE Evening needs 1 more")
        self.assertNotContains(dashboard, 'id="roomChart"')

    def test_section_occupancy_boundaries(self):
        for size in (0, 18, 20, 22, 23, 24, 25, 26, 49, 50, 51, 58, 70, 75, 76):
            with self.subTest(students=size):
                sections = section_occupancy(size)
                self.assertEqual(len(sections), (size + MAX_STUDENTS_PER_SECTION - 1) // MAX_STUDENTS_PER_SECTION)
                self.assertEqual(sum(section["students"] for section in sections), size)
                for section in sections:
                    self.assertLessEqual(section["students"], 25)
                    self.assertEqual(section["remaining_seats"], 25 - section["students"])
                    expected = "Full" if section["students"] == 25 else "Near Full" if section["students"] >= 23 else "Getting Full" if section["students"] >= 20 else "Normal"
                    self.assertEqual(section["status"], expected)
        self.assertEqual([section["students"] for section in section_occupancy(70)], [25, 25, 20])


class SharedCapacityAnalysisTests(TestCase):
    def setUp(self):
        self.department = Department.objects.create(code="IT", name="Information Technology")
        self.it = Major.objects.create(code="SOFT", name="Software Engineering", department=self.department)
        self.ir = Major.objects.create(code="IR", name="Information Research", department=self.department)

    def rooms(self, count):
        Room.objects.bulk_create([Room(code=f"R{index}", capacity=25, active=True) for index in range(count)])
        Room.objects.create(code="INACTIVE", capacity=25, active=False)

    def teachers(self, shift, count):
        Teacher.objects.bulk_create([
            Teacher(name=f"{shift} Teacher {index}", department=self.department, available_shifts=[shift])
            for index in range(count)
        ])

    def enroll(self, major, shift, count):
        Student.objects.bulk_create([
            Student(student_id=f"{major.code}-{shift}-{index}", student_name="Student", department=self.department, major=major, shift=shift)
            for index in range(count)
        ])

    def group(self, major, shift):
        return next(row for row in capacity_analysis()["groups"] if row["major_id"] == major.pk and row["shift"] == shift)

    def test_required_sections_and_remaining_seats_boundaries(self):
        for count in (0, 18, 24, 25, 26, 48, 50, 51, 63, 70, 75, 76):
            with self.subTest(students=count):
                required = (count + 24) // 25
                self.assertEqual(required_sections(count), required)
                self.assertEqual(sum(section["students"] for section in section_occupancy(count)), count)
                self.assertTrue(all(section["students"] <= 25 for section in section_occupancy(count)))
                self.assertEqual(required * 25 - count, sum(section["remaining_seats"] for section in section_occupancy(count)))

    def test_shared_rooms_used_and_room_capacity_limit(self):
        self.rooms(4)
        self.teachers(Shift.MORNING, 5)
        self.enroll(self.it, Shift.MORNING, 63)
        self.enroll(self.ir, Shift.MORNING, 18)
        analysis = capacity_analysis()
        morning = next(pool for pool in analysis["shift_pools"] if pool["shift"] == Shift.MORNING)
        self.assertEqual((morning["required"], morning["capacity"], morning["remaining"], morning["status"]), (4, 4, 0, "At Room Capacity"))
        it = self.group(self.it, Shift.MORNING)
        self.assertEqual((it["required"], it["required_rooms"], it["current_supported_capacity"], it["remaining_seats"]), (3, 3, 75, 12))
        self.assertEqual(it["status"], "Available")
        self.assertFalse(it["can_open_section"])
        self.assertEqual(it["shift_rooms_available"], 4)

        business = Department.objects.create(code="BUS", name="Business")
        business_major = Major.objects.create(code="BBA", name="Business", department=business)
        Teacher.objects.create(name="Business Teacher", department=business, available_shifts=[Shift.MORNING])
        Student.objects.bulk_create([
            Student(student_id=f"BUS-{index}", student_name="Student", department=business, major=business_major, shift=Shift.MORNING)
            for index in range(20)
        ])
        morning = next(pool for pool in capacity_analysis()["shift_pools"] if pool["shift"] == Shift.MORNING)
        self.assertEqual((morning["required"], morning["capacity"], morning["shortage"]), (5, 4, 1))
        self.assertEqual(self.group(self.it, Shift.MORNING)["status"], "Room Capacity Limit")

    def test_three_of_four_rooms_leaves_one_available(self):
        self.rooms(4)
        self.teachers(Shift.MORNING, 3)
        self.enroll(self.it, Shift.MORNING, 48)
        self.enroll(self.ir, Shift.MORNING, 18)
        morning = next(pool for pool in capacity_analysis()["shift_pools"] if pool["shift"] == Shift.MORNING)
        self.assertEqual((morning["required"], morning["remaining"], morning["status"]), (3, 1, "Available"))
        it = self.group(self.it, Shift.MORNING)
        self.assertEqual(it["remaining_seats"], 2)
        self.assertEqual(it["status"], "Near Capacity")  # No spare teacher despite the spare room.

    def test_another_section_requires_both_teacher_and_room(self):
        self.rooms(2)
        self.teachers(Shift.MORNING, 2)
        self.enroll(self.it, Shift.MORNING, 25)
        row = self.group(self.it, Shift.MORNING)
        self.assertEqual((row["remaining_seats"], row["potential_additional_sections"], row["supported_remaining_seats"]), (0, 1, 25))
        self.assertEqual(row["status"], "Available")
        self.assertTrue(row["can_open_section"])
        Student.objects.bulk_create([Student(student_id="NEXT", student_name="Next", department=self.department, major=self.it, shift=Shift.MORNING)])
        row = self.group(self.it, Shift.MORNING)
        self.assertEqual((row["required"], row["remaining_seats"], row["status"]), (2, 24, "Available"))
        self.assertEqual(row["sections"][0]["status"], "Full")

    def test_recommend_afternoon_for_near_capacity_morning(self):
        self.rooms(4)
        self.teachers(Shift.MORNING, 2)
        self.teachers(Shift.AFTERNOON, 2)
        self.teachers(Shift.EVENING, 1)
        self.enroll(self.it, Shift.MORNING, 48)
        self.enroll(self.it, Shift.AFTERNOON, 30)
        self.enroll(self.it, Shift.EVENING, 18)
        morning = self.group(self.it, Shift.MORNING)
        self.assertEqual((morning["status"], morning["remaining_seats"]), ("Near Capacity", 2))
        self.assertEqual((morning["recommended_shift"], morning["recommended_capacity"]), (Shift.AFTERNOON, 20))

    def test_near_capacity_is_five_or_fewer_supported_seats(self):
        self.rooms(1)
        self.teachers(Shift.MORNING, 1)
        self.enroll(self.it, Shift.MORNING, 19)
        self.assertEqual((self.group(self.it, Shift.MORNING)["supported_remaining_seats"], self.group(self.it, Shift.MORNING)["status"]), (6, "Available"))
        Student.objects.create(student_id="TWENTIETH", student_name="Student", department=self.department, major=self.it, shift=Shift.MORNING)
        self.assertEqual((self.group(self.it, Shift.MORNING)["supported_remaining_seats"], self.group(self.it, Shift.MORNING)["status"]), (5, "Near Capacity"))

    def test_empty_alternative_shift_is_considered_if_supported(self):
        self.rooms(1)
        self.teachers(Shift.MORNING, 1)
        self.teachers(Shift.AFTERNOON, 1)
        self.enroll(self.it, Shift.MORNING, 25)
        morning = self.group(self.it, Shift.MORNING)
        self.assertEqual(morning["status"], "Full")
        self.assertEqual((morning["recommended_shift"], morning["recommended_capacity"]), (Shift.AFTERNOON, 25))

    def test_full_morning_and_afternoon_recommend_evening(self):
        self.rooms(1)
        for shift in Shift.values:
            self.teachers(shift, 1)
        self.enroll(self.it, Shift.MORNING, 25)
        self.enroll(self.it, Shift.AFTERNOON, 25)
        self.enroll(self.it, Shift.EVENING, 18)
        morning = self.group(self.it, Shift.MORNING)
        self.assertEqual(morning["status"], "Full")
        self.assertEqual((morning["recommended_shift"], morning["recommended_capacity"]), (Shift.EVENING, 7))
        self.assertEqual(self.group(self.it, Shift.AFTERNOON)["status"], "Full")

    def test_no_alternative_when_all_shifts_full(self):
        self.rooms(1)
        for shift in Shift.values:
            self.teachers(shift, 1)
            self.enroll(self.it, shift, 25)
        for row in capacity_analysis()["groups"]:
            self.assertEqual(row["status"], "Full")
            self.assertIsNone(row["recommended_shift"])
            self.assertEqual(row["recommendation"], "No alternative shift currently has sufficient capacity.")

    def test_inactive_teacher_does_not_support_enrollment(self):
        self.rooms(1)
        Teacher.objects.create(name="Inactive", department=self.department, available_shifts=[Shift.MORNING], active=False)
        self.enroll(self.it, Shift.MORNING, 18)
        row = self.group(self.it, Shift.MORNING)
        self.assertEqual((row["teacher_capacity"], row["teacher_shortage"], row["status"]), (0, 1, "Teacher Shortage"))

    def test_no_enrollment_has_no_false_warnings(self):
        self.rooms(4)
        self.assertEqual(capacity_analysis()["groups"], [])
        self.assertTrue(all(pool["required"] == 0 for pool in capacity_analysis()["shift_pools"]))
        user = User.objects.create_user(email="empty@example.com", full_name="Empty Analyst", password="SafePass123!")
        self.client.force_login(user)
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, "No enrollment data yet.")
        self.assertNotContains(response, "⚠ Shortage")
        self.assertEqual(response.context["availability_alerts"], [])
