"""One database-backed workbook with native Excel charts and linked source cells."""

from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Attendance, Performance, Student
from .services import dashboard_context


BLUE = "2563EB"
INK = "111827"
MUTED = "6B7280"
PALE_BLUE = "E6EEFF"
SEPARATOR = Side(style="hair", color="D8D4C8")
PERCENT_FORMAT = "0.0%"


def _excel_value(value):
    """Imported text must not become an executable Excel formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _data_sheet(sheet, headers, rows, empty_message, percentage_columns=(), decimal_columns=()):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    widths = [len(header) for header in headers]
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        cell.font = Font(name="Calibri", size=11, bold=True, color=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=BLUE))
    sheet.row_dimensions[1].height = 30

    count = 0
    for row in rows:
        values = [_excel_value(value) for value in row]
        sheet.append(values)
        count += 1
        for column, value in enumerate(values, 1):
            widths[column - 1] = max(widths[column - 1], len(str(value)) if value is not None else 0)
            cell = sheet.cell(count + 1, column)
            cell.font = Font(name="Calibri", size=11, color=INK)
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left")
            cell.border = Border(bottom=SEPARATOR)
            if column in percentage_columns and value is not None:
                cell.number_format = PERCENT_FORMAT
            elif column in decimal_columns and value is not None:
                cell.number_format = "0.0"

    if count:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count + 1}"
    else:
        sheet.cell(2, 1, empty_message).font = Font(italic=True, color=MUTED)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 3, 14), 40)


def _dashboard_sheet(sheet, context, generated_at):
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:Q2")
    title = sheet["A1"]
    title.value = "UNIVERSITY DATA ANALYTICS REPORT"
    title.font = Font(name="Calibri", size=20, bold=True, color=INK)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 18
    sheet["A3"] = "Report generated"
    sheet["A3"].font = Font(color=MUTED, bold=True)
    sheet["C3"] = generated_at.replace(tzinfo=None)
    sheet["C3"].number_format = "dd mmm yyyy, hh:mm"

    kpis = context["kpis"]
    entries = [
        ("Total Students", kpis["students"], None),
        ("Active Teachers", kpis["teachers"], None),
        ("Active Rooms", kpis["rooms"], None),
        ("Departments", kpis["departments"], None),
        ("Majors", kpis["majors"], None),
        ("Average Score", kpis["average_score"], "0.0"),
        ("Pass Rate", kpis["pass_rate"] / 100 if kpis["pass_rate"] is not None else None, PERCENT_FORMAT),
        ("Average Attendance", kpis["average_attendance"] / 100 if kpis["average_attendance"] is not None else None, PERCENT_FORMAT),
        ("Attendance Risk Count", kpis["risk_count"], None),
    ]
    for row_number, (label, value, number_format) in enumerate(entries, 6):
        label_cell = sheet.cell(row_number, 1, label)
        label_cell.font = Font(name="Calibri", size=11, color=MUTED)
        label_cell.border = Border(bottom=SEPARATOR)
        value_cell = sheet.cell(row_number, 3, value if value is not None else "No Data")
        value_cell.font = Font(name="Calibri", size=13, bold=True, color=INK)
        value_cell.border = Border(bottom=SEPARATOR)
        if value is not None and number_format:
            value_cell.number_format = number_format

    if not context["has_data"]:
        sheet["A17"] = "No enrollment data available; enrollment charts are omitted."
    if not context["has_performance"]:
        sheet["A18"] = "No performance data available; performance charts are omitted."
    if not context["has_attendance"]:
        sheet["A19"] = "No attendance data available; attendance charts are omitted."
    for row in (17, 18, 19):
        if sheet.cell(row, 1).value:
            sheet.cell(row, 1).font = Font(italic=True, color=MUTED)

    sheet.column_dimensions["A"].width = 29
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 23
    for column in range(4, 19):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    sheet.freeze_panes = "A5"


def _chart_source(helper, index, title, labels, values):
    """Return cell references; every chart series reads its hidden worksheet cells."""
    column = 1 + index * 3
    helper.cell(1, column, title)
    helper.cell(1, column + 1, "Value")
    for row_number, (label, value) in enumerate(zip(labels, values), 2):
        helper.cell(row_number, column, _excel_value(label))
        helper.cell(row_number, column + 1, value)
    categories = Reference(helper, min_col=column, min_row=2, max_row=len(labels) + 1)
    data = Reference(helper, min_col=column + 1, min_row=1, max_row=len(labels) + 1)
    return categories, data


def _add_chart(dashboard, helper, index, title, labels, values, kind="bar", horizontal=False, colors=None):
    if not labels:
        return
    if len(labels) != len(values):
        raise ValueError(f"Chart data length mismatch: {title}")
    categories, data = _chart_source(helper, index, title, labels, values)
    chart = PieChart() if kind == "pie" else BarChart()
    chart.title = title
    chart.width = 16
    chart.height = min(13, max(8.5, 4 + len(labels) * 0.45)) if horizontal else 8.5
    chart.visible_cells_only = False  # Chart data lives on a hidden, editable helper sheet.
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    if kind == "pie":
        chart.legend.position = "b"
        chart.series[0].data_points = [
            DataPoint(idx=position, spPr=GraphicalProperties(solidFill=color))
            for position, color in enumerate(colors)
        ]
    else:
        chart.type = "bar" if horizontal else "col"
        chart.grouping = "clustered"
        chart.legend = None
        chart.series[0].graphicalProperties.solidFill = BLUE
    dashboard.add_chart(chart, f'{"A" if index % 2 == 0 else "J"}{21 + (index // 2) * 27}')


def build_analytics_workbook(generated_at=None):
    """Build a read-only snapshot from current ORM data and existing analytics."""
    generated_at = generated_at or timezone.localtime()
    if not isinstance(generated_at, datetime):
        raise TypeError("generated_at must be a datetime")
    context = dashboard_context()
    groups = context["capacity_analysis"]["groups"]
    group_by_major_shift = {(row["major_id"], row["shift"]): row for row in groups}
    room_by_shift = {row["shift"]: row for row in context["capacity_analysis"]["shift_pools"]}

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    enrollment = workbook.create_sheet("Enrollment")
    performance = workbook.create_sheet("Performance")
    attendance = workbook.create_sheet("Attendance")
    capacity = workbook.create_sheet("Capacity")
    helper = workbook.create_sheet("_ChartData")
    _dashboard_sheet(dashboard, context, generated_at)

    _data_sheet(
        enrollment,
        ["Student ID", "Student Name", "Department", "Major", "Shift", "Group Required Sections", "Group Capacity Status"],
        ([student.student_id, student.student_name, student.department.code, student.major.name, student.shift,
          group_by_major_shift[(student.major_id, student.shift)]["required"],
          group_by_major_shift[(student.major_id, student.shift)]["status"]]
         for student in Student.objects.select_related("department", "major").order_by("student_id")),
        "No enrollment data available",
    )
    _data_sheet(
        performance,
        ["Student ID", "Student Name", "Subject Code", "Subject Name", "Score", "Grade", "Pass/Fail", "Grade Point"],
        ([record.student.student_id, record.student.student_name, record.subject.code, record.subject.name,
          float(record.score), record.letter_grade, "Pass" if record.passed else "Fail", float(record.grade_point)]
         for record in Performance.objects.select_related("student", "subject").order_by("student__student_id", "subject__code")),
        "No performance data available",
        decimal_columns=(5, 8),
    )
    _data_sheet(
        attendance,
        ["Student ID", "Student Name", "Subject Code", "Subject Name", "Total Sessions", "Absent Count", "Attendance %", "Attendance Status"],
        ([record.student.student_id, record.student.student_name, record.subject.code, record.subject.name,
          record.total_sessions, record.absent_count, float(record.attendance_percentage) / 100, record.status]
         for record in Attendance.objects.select_related("student", "subject").order_by("student__student_id", "subject__code")),
        "No attendance data available",
        percentage_columns=(7,),
    )
    _data_sheet(
        capacity,
        ["Department", "Major", "Shift", "Students", "Required Sections", "Teacher Capacity",
         "Required Rooms", "Active Rooms", "Rooms Remaining", "Remaining Supported Seats",
         "Capacity Status", "Room Capacity Status", "Recommended Alternative Shift"],
        ([row["major__department__code"], row["major__name"], row["shift"], row["students"],
          row["required"], row["teacher_capacity"], row["required_rooms"], row["shift_rooms_available"],
          row["shift_rooms_remaining"], row["supported_remaining_seats"], row["status"],
          room_by_shift[row["shift"]]["status"], row["recommended_shift"] or ""] for row in groups),
        "No enrollment data available for capacity analysis",
    )

    required_labels = [f'{row["major__department__code"]} / {row["major__name"]} · {row["shift"]}' for row in groups]
    chart_specs = [
        ("Enrollment by Department", context["enrollment_department"]["labels"],
         [int(value) for value in context["enrollment_department"]["values"]], "bar", False, None),
        ("Enrollment by Major", context["enrollment_major"]["labels"],
         [int(value) for value in context["enrollment_major"]["values"]], "bar", True, None),
        ("Students by Shift", context["enrollment_shift"]["labels"],
         [int(value) for value in context["enrollment_shift"]["values"]], "bar", False, None),
        ("Required Sections by Major / Shift", required_labels,
         [row["required"] for row in groups], "bar", True, None),
        ("Average Score by Subject", context["subject_scores"]["labels"],
         context["subject_scores"]["values"], "bar", True, None),
        ("Pass vs Fail", context["pass_fail"]["labels"] if context["has_performance"] else [],
         [int(value) for value in context["pass_fail"]["values"]], "pie", False, ["16A34A", "DC2626"]),
        ("Attendance Risk Distribution", context["attendance_status"]["labels"] if context["has_attendance"] else [],
         [int(value) for value in context["attendance_status"]["values"]], "pie", False,
         ["16A34A", "D97706", "EA580C", "DC2626"]),
    ]
    chart_index = 0
    for title, labels, values, kind, horizontal, colors in chart_specs:
        if labels:
            _add_chart(dashboard, helper, chart_index, title, labels, values, kind, horizontal, colors)
            chart_index += 1
    helper.sheet_state = "hidden"
    return workbook
