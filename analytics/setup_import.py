"""Preview-first, create-only imports for individual setup datasets."""

from io import BytesIO

import pandas as pd
from django.db import IntegrityError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .demo_workbooks import MASTER_ROWS
from .forms import TeacherForm
from .master_import import active_boolean, positive_integer
from .models import Department, Major, Room, Shift, Subject, Teacher
from .services import clean_cell


SETUP_SPECS = {
    "department": ["code", "name"],
    "major": ["code", "name", "department_code"],
    "subject": ["code", "name", "major_code"],
    "teacher": ["name", "email", "department_code", "available_shifts", "subjects_can_teach", "active"],
    "room": ["code", "capacity", "active"],
}


def setup_template(kind):
    headers = SETUP_SPECS[kind]
    source_rows = MASTER_ROWS[kind.title() + "s"][1:]
    examples = [list(row) for row in source_rows]
    book = Workbook()
    sheet = book.active
    sheet.title = kind.title() + "s"
    sheet.append(headers)
    for row in examples:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="111827")
        cell.fill = PatternFill("solid", fgColor="E6EEFF")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(38, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
    sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def validate_setup_workbook(file, kind):
    headers = SETUP_SPECS[kind]
    try:
        frame = pd.read_excel(file, engine="openpyxl", dtype=object)
    except Exception as exc:
        raise ValueError(f"Excel file could not be read: {exc}") from exc
    frame.columns = [str(column).strip().lower() for column in frame.columns]
    if len(frame.columns) != len(set(frame.columns)):
        raise ValueError("Duplicate column headers are not allowed")
    missing_columns = [column for column in headers if column not in frame.columns]
    if missing_columns:
        raise ValueError("Missing required columns: " + ", ".join(missing_columns))
    if frame.empty:
        raise ValueError("The workbook has no data rows")

    departments = {item.code.upper(): item for item in Department.objects.all()}
    majors = {item.code.upper(): item for item in Major.objects.select_related("department")}
    subjects = {item.code.upper(): item for item in Subject.objects.select_related("major")}
    existing_codes = {
        "department": set(departments), "major": set(majors), "subject": set(subjects),
        "room": {room.code.upper() for room in Room.objects.all()},
    }
    teacher_emails = {item.email.lower() for item in Teacher.objects.exclude(email="")}
    teacher_names = {(item.name.casefold(), item.department.code.upper()) for item in Teacher.objects.select_related("department")}
    summary = {"total_rows": len(frame), "valid_rows": 0, "rejected_rows": 0,
               "duplicate_rows": 0, "missing_value_rows": 0, "invalid_references": 0,
               "invalid_numeric_rows": 0}
    valid_rows, issues, seen = [], [], set()

    for index, raw in frame.iterrows():
        row_number = index + 2
        row = {column: clean_cell(raw.get(column)) for column in headers}
        errors = []
        required = [column for column in headers if column not in {"email", "subjects_can_teach"}]
        missing = [column for column in required if not row[column]]
        if missing:
            summary["missing_value_rows"] += 1
            errors.append("Missing: " + ", ".join(missing))
        if kind == "teacher":
            key = (row["name"].casefold(), row["department_code"].upper())
            identities = {("name", key)}
            if row["email"]:
                identities.add(("email", row["email"].lower()))
            exists = key in teacher_names or (row["email"] and row["email"].lower() in teacher_emails)
        else:
            row["code"] = row["code"].upper()
            identities = {row["code"]} if row["code"] else set()
            exists = row["code"] in existing_codes[kind]
        if identities & seen or exists:
            summary["duplicate_rows"] += 1
            if identities & seen:
                errors.append("Duplicate record in workbook")
        if exists:
            errors.append("Record already exists in the database; edit it manually")
        seen.update(identities)

        if kind in {"department", "major", "subject", "room"}:
            if len(row["code"]) > (30 if kind == "room" else 20):
                errors.append("Code is too long")
        if kind in {"department", "major", "subject", "teacher"} and len(row["name"]) > 150:
            errors.append("Name is too long")
        if kind == "major":
            row["department_code"] = row["department_code"].upper()
            if row["department_code"] and row["department_code"] not in departments:
                summary["invalid_references"] += 1
                errors.append("Unknown department_code; import Departments first")
        elif kind == "subject":
            row["major_code"] = row["major_code"].upper()
            if row["major_code"] and row["major_code"] not in majors:
                summary["invalid_references"] += 1
                errors.append("Unknown major_code; import Majors first")
        elif kind == "room":
            try:
                row["capacity"] = positive_integer(row["capacity"])
            except ValueError as exc:
                summary["invalid_numeric_rows"] += 1
                errors.append(f"capacity: {exc}")
            try:
                row["active"] = active_boolean(row["active"])
            except ValueError as exc:
                errors.append(str(exc))
        elif kind == "teacher":
            row["department_code"] = row["department_code"].upper()
            department = departments.get(row["department_code"])
            if row["department_code"] and not department:
                summary["invalid_references"] += 1
                errors.append("Unknown department_code; import Departments first")
            shifts = [value.strip() for value in row["available_shifts"].split(",") if value.strip()]
            shift_lookup = {value.lower(): value for value in Shift.values}
            if len(shifts) != len(set(value.lower() for value in shifts)) or any(value.lower() not in shift_lookup for value in shifts):
                errors.append("Available shifts must be unique Morning, Afternoon, or Evening values")
            row["available_shifts"] = [shift_lookup[value.lower()] for value in shifts if value.lower() in shift_lookup]
            codes = [value.strip().upper() for value in row["subjects_can_teach"].split(",") if value.strip()]
            if len(codes) != len(set(codes)):
                errors.append("Duplicate subject code in subjects_can_teach")
            unknown = [code for code in codes if code not in subjects]
            if unknown:
                summary["invalid_references"] += 1
                errors.append("Unknown subject code(s): " + ", ".join(unknown))
            row["subjects_can_teach"] = codes
            try:
                row["active"] = active_boolean(row["active"])
            except ValueError as exc:
                errors.append(str(exc))
            if department and not unknown:
                form = TeacherForm(data={"name": row["name"], "email": row["email"],
                    "department": department.pk, "available_shifts": row["available_shifts"],
                    "subjects_can_teach": [subjects[code].pk for code in codes],
                    "active": "on" if row["active"] is True else ""})
                if not form.is_valid():
                    errors.extend(f"{field}: {', '.join(messages)}" for field, messages in form.errors.items())
        if errors:
            summary["rejected_rows"] += 1
            issues.append({"row": row_number, "errors": errors, "data": {key: clean_cell(raw.get(key)) for key in headers}})
        else:
            summary["valid_rows"] += 1
            valid_rows.append(row)
    return summary, valid_rows, issues


def import_setup_rows(kind, rows):
    """Create validated rows only; concurrent changes abort the surrounding transaction."""
    def resolve_code(model, code):
        matches = list(model.objects.filter(code__iexact=code)[:2])
        if len(matches) != 1:
            raise ValueError(f"{model.__name__} code {code} was removed or became ambiguous after validation")
        return matches[0]

    created = 0
    for row in rows:
        if kind == "teacher":
            department = resolve_code(Department, row["department_code"])
            if Teacher.objects.filter(name__iexact=row["name"], department=department).exists() or (
                row["email"] and Teacher.objects.filter(email__iexact=row["email"]).exists()
            ):
                raise ValueError(f"Teacher {row['name']} already exists")
            subjects = [resolve_code(Subject, code) for code in row["subjects_can_teach"]]
            form = TeacherForm(data={"name": row["name"], "email": row["email"],
                "department": department.pk, "available_shifts": row["available_shifts"],
                "subjects_can_teach": [subject.pk for subject in subjects],
                "active": "on" if row["active"] else ""})
            if not form.is_valid():
                raise ValueError(f"Teacher {row['name']} is no longer valid: {form.errors.as_text()}")
            teacher = Teacher.objects.create(name=row["name"], email=row["email"], department=department,
                available_shifts=row["available_shifts"], active=row["active"])
            teacher.subjects_can_teach.set(subjects)
        else:
            model = {"department": Department, "major": Major, "subject": Subject, "room": Room}[kind]
            if model.objects.filter(code__iexact=row["code"]).exists():
                raise ValueError(f"{kind.title()} {row['code']} already exists")
            if kind == "department":
                Department.objects.create(code=row["code"], name=row["name"])
            elif kind == "major":
                Major.objects.create(code=row["code"], name=row["name"],
                                     department=resolve_code(Department, row["department_code"]))
            elif kind == "subject":
                Subject.objects.create(code=row["code"], name=row["name"],
                                       major=resolve_code(Major, row["major_code"]))
            else:
                Room.objects.create(code=row["code"], capacity=row["capacity"], active=row["active"])
        created += 1
    return created, 0
