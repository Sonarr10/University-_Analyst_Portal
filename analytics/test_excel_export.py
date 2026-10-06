from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .capacity import capacity_analysis
from .excel_export import build_analytics_workbook
from .models import Attendance, Department, Major, Performance, Room, Shift, Student, Subject, Teacher, User
from .services import dashboard_context


class UniversityExcelExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email="report@example.com", full_name="Report Analyst", password="SafePass123!"
        )
        self.department = Department.objects.create(code="CSE", name="Computer Science")
        self.major = Major.objects.create(code="SE", name="Software Engineering", department=self.department)
        self.subject = Subject.objects.create(code="DB101", name="Database Systems", major=self.major)
        Teacher.objects.create(name="Morning Teacher", department=self.department, available_shifts=[Shift.MORNING, Shift.AFTERNOON])
        Teacher.objects.create(name="Inactive Teacher", department=self.department, available_shifts=[Shift.MORNING], active=False)
        Room.objects.create(code="A101", capacity=25, active=True)
        Room.objects.create(code="A102", capacity=25, active=False)

    def add_student(self, number, shift=Shift.MORNING, name="Student"):
        return Student.objects.create(
            student_id=f"ST{number:03}", student_name=name,
            department=self.department, major=self.major, shift=shift,
        )

    def workbook_from_response(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("export_university_analytics"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("University_Analytics_Report_", response["Content-Disposition"])
        return response, load_workbook(BytesIO(response.content))

    def test_empty_database_report_has_real_setup_counts_and_no_broken_charts(self):
        response, workbook = self.workbook_from_response()
        self.assertEqual(workbook.sheetnames, [
            "Dashboard", "Enrollment", "Performance", "Attendance", "Capacity", "_ChartData"
        ])
        dashboard = workbook["Dashboard"]
        self.assertEqual(dashboard["C6"].value, 0)
        self.assertEqual(dashboard["C7"].value, 1)  # Only the active teacher counts.
        self.assertEqual(dashboard["C8"].value, 1)  # Only the active room counts.
        self.assertEqual(dashboard["C11"].value, "No Data")
        self.assertEqual(len(dashboard._charts), 0)
        self.assertEqual(workbook["_ChartData"].sheet_state, "hidden")
        for name in ("Enrollment", "Performance", "Attendance", "Capacity"):
            self.assertIn("No ", workbook[name]["A2"].value)
            self.assertEqual(workbook[name].freeze_panes, "A2")
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertFalse(any(name.startswith("xl/charts/") for name in archive.namelist()))

    def test_full_report_uses_native_cell_linked_charts_and_numeric_values(self):
        first = self.add_student(1, name="=SUM(1,1)")
        second = self.add_student(2)
        Performance.objects.create(student=first, subject=self.subject, score=75, letter_grade="B", grade_point=3, passed=True)
        Performance.objects.create(student=second, subject=self.subject, score=40, letter_grade="F", grade_point=0, passed=False)
        Attendance.objects.create(student=first, subject=self.subject, total_sessions=20, absent_count=5,
                                  attendance_percentage=75, status="Warning")
        Attendance.objects.create(student=second, subject=self.subject, total_sessions=20, absent_count=10,
                                  attendance_percentage=50, status="Critical")
        context = dashboard_context()
        analysis = capacity_analysis()
        response, workbook = self.workbook_from_response()
        dashboard = workbook["Dashboard"]
        self.assertEqual(len(dashboard._charts), 7)
        self.assertEqual(dashboard["C6"].value, context["kpis"]["students"])
        self.assertEqual(dashboard["C11"].value, context["kpis"]["average_score"])
        self.assertAlmostEqual(dashboard["C12"].value, context["kpis"]["pass_rate"] / 100)
        self.assertAlmostEqual(dashboard["C13"].value, context["kpis"]["average_attendance"] / 100)
        self.assertEqual(dashboard["C12"].number_format, "0.0%")
        self.assertEqual(dashboard["C14"].value, context["kpis"]["risk_count"])
        self.assertIsInstance(dashboard["C3"].value, datetime)

        enrollment = workbook["Enrollment"]
        self.assertEqual(enrollment.max_row, 3)
        self.assertEqual(enrollment["B2"].value, "'=SUM(1,1)")
        self.assertEqual(enrollment["F2"].value, analysis["groups"][0]["required"])
        self.assertEqual(enrollment["G2"].value, analysis["groups"][0]["status"])
        performance = workbook["Performance"]
        self.assertEqual(performance["E2"].value, 75)
        self.assertEqual(performance["E2"].data_type, "n")
        self.assertEqual(performance["H2"].value, 3)
        attendance = workbook["Attendance"]
        self.assertAlmostEqual(attendance["G2"].value, 0.75)
        self.assertEqual(attendance["G2"].data_type, "n")
        self.assertEqual(attendance["G2"].number_format, "0.0%")
        capacity = workbook["Capacity"]
        row = analysis["groups"][0]
        self.assertEqual([capacity.cell(2, col).value for col in (4, 5, 6, 7, 8, 10, 11)], [
            row["students"], row["required"], row["teacher_capacity"], row["required_rooms"],
            row["shift_rooms_available"], row["supported_remaining_seats"], row["status"],
        ])

        helper = workbook["_ChartData"]
        self.assertEqual(helper.sheet_state, "hidden")
        self.assertEqual(helper["A2"].value, "CSE")
        self.assertEqual(helper["B2"].value, 2)
        self.assertEqual(helper["N2"].value, 57.5)  # Real subject average, not a chart screenshot.
        self.assertEqual(helper["Q2"].value, 1)
        self.assertEqual(helper["Q3"].value, 1)
        self.assertEqual([helper[f"T{row}"].value for row in range(2, 6)], [0, 1, 0, 1])
        for chart in dashboard._charts:
            self.assertFalse(chart.visible_cells_only)
            self.assertIn("'_ChartData'!", chart.series[0].val.numRef.f)
            self.assertIn("'_ChartData'!", chart.series[0].cat.numRef.f)
        self.assertEqual(dashboard._charts[0].series[0].graphicalProperties.solidFill.srgbClr, "2563EB")
        self.assertEqual(
            [point.spPr.solidFill.srgbClr for point in dashboard._charts[5].series[0].data_points],
            ["16A34A", "DC2626"],
        )
        with ZipFile(BytesIO(response.content)) as archive:
            names = archive.namelist()
            self.assertEqual(len([name for name in names if name.startswith("xl/charts/chart") and name.endswith(".xml")]), 7)
            self.assertFalse(any(name.startswith("xl/media/") for name in names))
            chart_xml = "".join(archive.read(name).decode() for name in names if name.startswith("xl/charts/chart") and name.endswith(".xml"))
            self.assertIn("_ChartData", chart_xml)
            self.assertNotIn("#REF!", chart_xml)

    def test_missing_attendance_skips_only_attendance_chart(self):
        student = self.add_student(1)
        Performance.objects.create(student=student, subject=self.subject, score=60,
                                   letter_grade="C", grade_point=2, passed=True)
        _, workbook = self.workbook_from_response()
        self.assertEqual(len(workbook["Dashboard"]._charts), 6)
        self.assertIn("No attendance data available", workbook["Attendance"]["A2"].value)
        self.assertEqual(workbook["Dashboard"]["C13"].value, "No Data")

    def test_missing_performance_skips_only_performance_charts(self):
        student = self.add_student(1)
        Attendance.objects.create(student=student, subject=self.subject, total_sessions=20,
                                  absent_count=0, attendance_percentage=100, status="Good")
        _, workbook = self.workbook_from_response()
        self.assertEqual(len(workbook["Dashboard"]._charts), 5)
        self.assertIn("No performance data available", workbook["Performance"]["A2"].value)
        self.assertEqual(workbook["Dashboard"]["C11"].value, "No Data")
        self.assertEqual(workbook["Dashboard"]["C13"].value, 1)

    def test_capacity_recommendation_matches_website_analysis(self):
        Student.objects.bulk_create([
            Student(student_id=f"ST{index:03}", student_name="Student", department=self.department,
                    major=self.major, shift=Shift.MORNING)
            for index in range(25)
        ])
        expected = capacity_analysis()["groups"][0]
        self.assertEqual(expected["recommended_shift"], Shift.AFTERNOON)
        _, workbook = self.workbook_from_response()
        capacity = workbook["Capacity"]
        self.assertEqual(capacity["E2"].value, expected["required"])
        self.assertEqual(capacity["F2"].value, expected["teacher_capacity"])
        self.assertEqual(capacity["H2"].value, expected["shift_rooms_available"])
        self.assertEqual(capacity["K2"].value, expected["status"])
        self.assertEqual(capacity["M2"].value, expected["recommended_shift"])

    def test_report_requires_login_and_button_is_on_dashboard(self):
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse("dashboard")), "Export Analytics Excel")
        self.client.logout()
        self.assertRedirects(
            self.client.get(reverse("export_university_analytics")),
            f"{reverse('login')}?next={reverse('export_university_analytics')}",
        )
